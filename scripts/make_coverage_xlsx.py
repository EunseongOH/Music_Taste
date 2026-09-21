# 아티스트 확보 현황 스프레드시트를 만든다 (행 배경색으로 구분).
#
#   1 트랙리스트 전체   연한 파랑
#   2 앨범 목록만       연한 초록
#   3 트랙리스트 일부   연한 노랑
#   4 Spotify 필요      연한 빨강
#
# 사용: python scripts/make_coverage_xlsx.py <입력.json> <출력.xlsx>

import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = sys.argv[1] if len(sys.argv) > 1 else "C:/Users/User/sortify-exports/아티스트_확보현황.json"
DST = sys.argv[2] if len(sys.argv) > 2 else "C:/Users/User/sortify-exports/아티스트_확보현황.xlsx"

LABEL = {
    1: "트랙리스트 전체 보유",
    2: "앨범 목록만 보유",
    3: "트랙리스트 일부 보유",
    4: "Spotify API 필요",
}
FILL = {
    1: "DDEBF7",   # 연한 파랑
    2: "E2EFDA",   # 연한 초록
    3: "FFF2CC",   # 연한 노랑
    4: "FCE4E4",   # 연한 빨강
}
NOTE = {
    1: "아는 앨범을 빠짐없이 곡까지 갖고 있다. Spotify 없이 소트를 끝까지 할 수 있다.",
    2: "앨범 이름은 아는데 곡이 없다. 앨범을 누르면 Spotify 를 부른다.",
    3: "곡이 있는 앨범도 있고, 아직 못 받은 앨범도 있다.",
    4: "연결이 없거나 믿을 수 없어 자체 DB 로는 아무것도 못 낸다. 전부 Spotify 를 부른다.",
}

rows = json.load(open(SRC, encoding="utf-8"))
if not rows:
    raise SystemExit("데이터가 비었다")

cols = ["구분"] + [k for k in rows[0].keys() if k != "구분"]
wb = Workbook()

# --- 안내 시트 ---
info = wb.active
info.title = "안내"
info["A1"] = "아티스트 확보 현황"
info["A1"].font = Font(bold=True, size=14)
info["A2"] = "행 배경색이 구분이다. 숫자는 지금 자체 DB 로 낼 수 있는 양이고, Spotify 는 한 번도 부르지 않고 집계했다."
for i, k in enumerate([1, 2, 3, 4], start=4):
    n = sum(1 for r in rows if r["구분"] == k)
    info.cell(row=i, column=1, value=LABEL[k]).fill = PatternFill("solid", fgColor=FILL[k])
    info.cell(row=i, column=2, value=f"{n}팀").fill = PatternFill("solid", fgColor=FILL[k])
    info.cell(row=i, column=3, value=NOTE[k]).fill = PatternFill("solid", fgColor=FILL[k])
info.cell(row=9, column=1, value="합계")
info.cell(row=9, column=2, value=f"{len(rows)}팀")
info.column_dimensions["A"].width = 24
info.column_dimensions["B"].width = 10
info.column_dimensions["C"].width = 78

# --- 목록 시트 ---
ws = wb.create_sheet("아티스트")
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

ws.append([("구분" if c == "구분" else c) for c in cols])
head_fill = PatternFill("solid", fgColor="44546A")
for j in range(1, len(cols) + 1):
    c = ws.cell(row=1, column=j)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = head_fill
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border

for r in rows:
    vals = []
    for c in cols:
        vals.append(LABEL[r["구분"]] if c == "구분" else r.get(c, ""))
    ws.append(vals)
    fill = PatternFill("solid", fgColor=FILL[r["구분"]])
    for j in range(1, len(cols) + 1):
        cell = ws.cell(row=ws.max_row, column=j)
        cell.fill = fill
        cell.border = border

widths = {"구분": 20, "아티스트": 28, "국가": 6, "연결근거": 11, "우선순위": 26, "spotify_id": 24}
for j, c in enumerate(cols, start=1):
    ws.column_dimensions[get_column_letter(j)].width = widths.get(c, 15)
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

wb.save(DST)
print(f"{DST} · {len(rows)}행")
for k in (1, 2, 3, 4):
    print(f"  {LABEL[k]}: {sum(1 for r in rows if r['구분'] == k)}팀")
